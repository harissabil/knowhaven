"""Document parsers (design doc 15.5): Markdown and DOCX per heading, PDF per page.
Also used for Confluence pages (converted to Markdown) so citations look the same."""

import io
import re

from brain.ingestion.chunking import heading_anchor, markdown_sections
from brain.ingestion.types import Section

_ALREADY_PREFIXED = re.compile(r"^\s*(?:(?:\d+(?:\.\d+)*|[a-zA-Z])[.)]|[-*•])\s+")
# In PDF extracted text, consecutive numbered items are sometimes joined on one line.
_INLINE_NUMBERED = re.compile(r"(?<=[.!?;:])\s+(?=\d{1,2}\.\s+[A-Z0-9])")


def parse_markdown(text: str) -> tuple[str | None, list[Section]]:
    return markdown_sections(text)


def _docx_numbering_formats(doc) -> dict[str, dict[str, str]]:
    """Map numId -> {ilvl -> numFmt} from word/numbering.xml if present."""
    try:
        numbering_part = doc.part.numbering_part
        if numbering_part is None or numbering_part.element is None:
            return {}
        root = numbering_part.element
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        abstract_fmts: dict[str, dict[str, str]] = {}
        for abs_num in root.findall("w:abstractNum", ns):
            abs_id = abs_num.get(f"{{{ns['w']}}}abstractNumId")
            if not abs_id:
                continue
            lvls: dict[str, str] = {}
            for lvl in abs_num.findall("w:lvl", ns):
                ilvl = lvl.get(f"{{{ns['w']}}}ilvl") or "0"
                fmt_el = lvl.find("w:numFmt", ns)
                fmt = fmt_el.get(f"{{{ns['w']}}}val") if fmt_el is not None else "decimal"
                lvls[ilvl] = fmt or "decimal"
            abstract_fmts[abs_id] = lvls
        num_to_fmt: dict[str, dict[str, str]] = {}
        for num in root.findall("w:num", ns):
            num_id = num.get(f"{{{ns['w']}}}numId")
            abs_el = num.find("w:abstractNumId", ns)
            abs_id = abs_el.get(f"{{{ns['w']}}}val") if abs_el is not None else None
            if num_id and abs_id and abs_id in abstract_fmts:
                num_to_fmt[num_id] = dict(abstract_fmts[abs_id])
        return num_to_fmt
    except Exception:  # noqa: BLE001
        return {}


def _docx_list_prefix(
    p, counters: dict[tuple[str, str], int], num_fmts: dict[str, dict[str, str]]
) -> str:
    """Return '1. ', '2. ', or '- ' for Word list paragraphs whose text lacks an explicit prefix."""
    style = (p.style.name or "").lower() if p.style is not None else ""
    pPr = getattr(p._p, "pPr", None)
    numPr = getattr(pPr, "numPr", None) if pPr is not None else None
    num_id: str | None = None
    ilvl = "0"
    if numPr is not None:
        if getattr(numPr, "numId", None) is not None and numPr.numId.val is not None:
            num_id = str(numPr.numId.val)
        if getattr(numPr, "ilvl", None) is not None and numPr.ilvl.val is not None:
            ilvl = str(numPr.ilvl.val)

    if num_id == "0":  # numId=0 in OOXML disables numbering on that paragraph
        num_id = None

    is_bullet_style = "bullet" in style
    is_number_style = "number" in style

    if num_id is not None:
        fmt = num_fmts.get(num_id, {}).get(ilvl, "bullet" if is_bullet_style else "decimal")
        if fmt == "bullet" or is_bullet_style:
            return "- "
        key = (num_id, ilvl)
        counters[key] = counters.get(key, 0) + 1
        return f"{counters[key]}. "

    if is_bullet_style:
        return "- "
    if is_number_style:
        key = ("style_num", ilvl)
        counters[key] = counters.get(key, 0) + 1
        return f"{counters[key]}. "
    return ""


def parse_docx(data: bytes) -> tuple[str | None, list[Section]]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    title = (doc.core_properties.title or "").strip() or None
    sections: list[Section] = []
    heading: str | None = None
    buf: list[str] = []
    num_fmts = _docx_numbering_formats(doc)
    counters: dict[tuple[str, str], int] = {}

    def flush():
        body = "\n".join(buf).strip()
        if body:
            sections.append(Section(heading_anchor(heading) if heading else None, heading, body))

    for p in doc.paragraphs:
        style = (p.style.name or "").lower() if p.style is not None else ""
        text = p.text.strip()
        if style == "title" and text:
            title = title or text
            counters.clear()
        elif style.startswith("heading") and text:
            counters.clear()
            if title is None and style == "heading 1" and not sections and not buf:
                title = text
                continue
            flush()
            heading, buf = text, []
        elif text:
            prefix = (
                "" if _ALREADY_PREFIXED.match(text) else _docx_list_prefix(p, counters, num_fmts)
            )
            if not prefix and "list" not in style:
                counters.pop(("style_num", "0"), None)
            buf.append(f"{prefix}{text}")
    flush()
    for table in doc.tables:  # tables (e.g. grade gaji) become one section each
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if rows:
            sections.append(Section(f"tabel-{len(sections) + 1}", "Tabel", "\n".join(rows)))
    return title, sections


def parse_pdf(data: bytes) -> tuple[str | None, list[Section]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    title = None
    if reader.metadata and reader.metadata.title:
        title = str(reader.metadata.title).strip() or None
    sections = []
    for i, page in enumerate(reader.pages, start=1):
        text = re.sub(r"[ \t]+\n", "\n", page.extract_text() or "").strip()
        text = _INLINE_NUMBERED.sub("\n", text)
        if text:
            sections.append(Section(f"p{i}", f"Halaman {i}", text))
    return title, sections


def parse_document(filename: str, data: bytes) -> tuple[str | None, list[Section]]:
    name = filename.lower()
    if name.endswith((".md", ".markdown", ".txt")):
        return parse_markdown(data.decode("utf-8", errors="replace"))
    if name.endswith(".docx"):
        return parse_docx(data)
    if name.endswith(".pdf"):
        return parse_pdf(data)
    raise ValueError("unsupported_file_type")
